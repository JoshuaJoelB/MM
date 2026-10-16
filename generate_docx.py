# -*- coding: utf-8 -*-
"""Generate Match Monster project documentation (DOCX)."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x1E, 0x90, 0xFF)   # bright blue
DARK   = RGBColor(0x0A, 0x14, 0x28)   # dark navy
GOLD   = RGBColor(0xC0, 0x8A, 0x00)   # muted gold
GREY   = RGBColor(0x55, 0x5F, 0x6B)

doc = Document()

# ------------------------------------------------------------------
# Base styles
# ------------------------------------------------------------------
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x21, 0x27, 0x2E)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, bold in [
    ("Heading 1", 16, ACCENT, True),
    ("Heading 2", 13, DARK, True),
    ("Heading 3", 11.5, GOLD, True),
]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = bold
    st.font.color.rgb = color
    # ensure east-asian fallback font set too
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "Calibri")

def p(text="", style=None, size=None, bold=False, color=None, align=None,
      space_after=None, italic=False):
    par = doc.add_paragraph(style=style)
    if align is not None:
        par.alignment = align
    if space_after is not None:
        par.paragraph_format.space_after = Pt(space_after)
    if text:
        run = par.add_run(text)
        run.bold = bold
        run.italic = italic
        if size: run.font.size = Pt(size)
        if color: run.font.color.rgb = color
    return par

def bullet(text, level=0, bold_prefix=None):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    par = doc.add_paragraph(style=style)
    if bold_prefix:
        r = par.add_run(bold_prefix)
        r.bold = True
    par.add_run(text)
    return par

def numbered(text):
    return doc.add_paragraph(text, style="List Number")

def h1(t): return p(t, style="Heading 1")
def h2(t): return p(t, style="Heading 2")
def h3(t): return p(t, style="Heading 3")

def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)

def style_table_cells(table, header_bg="1E90FF", header_fg=RGBColor(0xFF, 0xFF, 0xFF)):
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            for par in cell.paragraphs:
                par.paragraph_format.space_after = Pt(2)
                for run in par.runs:
                    run.font.size = Pt(10)
            if i == 0:
                # first column header-ish
                pass
    # header row
    for cell in table.rows[0].cells:
        set_cell_bg(cell, header_bg)
        for par in cell.paragraphs:
            for run in par.runs:
                run.font.bold = True
                run.font.color.rgb = header_fg

def make_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = htext
    for r in rows:
        cells = table.add_row().cells
        for i, val in enumerate(r):
            cells[i].text = str(val)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Inches(w)
    style_table_cells(table)
    return table

# ==================================================================
# TITLE PAGE
# ==================================================================
for _ in range(4):
    doc.add_paragraph()

p("MATCH MONSTER", size=34, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER)
p("An Educational Memory-Matching Web Game for Kids", size=14, italic=True,
  color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
p("Project Documentation", size=12, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_paragraph()
p("Project by: BSEMC 3rd Year", size=12, bold=True, color=DARK,
  align=WD_ALIGN_PARAGRAPH.CENTER)
p("Web-based · Puzzle · Educational · School Project", size=11, color=GREY,
  align=WD_ALIGN_PARAGRAPH.CENTER)

doc.add_page_break()

# ==================================================================
# 1. OVERVIEW
# ==================================================================
h1("1. Overview")
p("Match Monster is a kid-friendly, web-based memory-matching card game that turns "
  "school lessons into play. A player flips pairs of cards and matches the related "
  "items — images, words, and concepts — to complete a level. The game blends the "
  "classic \u201cconcentration\u201d mechanic with an educational curriculum across three "
  "subjects, and it wraps everything in a colorful, monster-themed world.")
p("It is a purely client-side application: there is no server, no account system, and "
  "no internet data collection. All progress, scores, and settings are stored locally "
  "on the player's own device, making it both simple to run and privacy-friendly.")

h2("1.1 Project Classification")
bullet("Platform: Web-based application (runs in any modern browser).")
bullet("Genre: Puzzle / memory (educational).")
bullet("Purpose: School project for BSEMC (Bachelor of Science in Entertainment and "
       "Multimedia Computing) 3rd Year.")
bullet("Audience: Children / young learners.")

h2("1.2 Key Features at a Glance")
bullet("A nickname-based login with per-player saved progress.")
bullet("Three playable subjects, each with a 10-level island map.")
bullet("Classic flip-and-match memory gameplay with a countdown timer.")
bullet("A 1\u20133 star rating and a points system that gates level unlocks.")
bullet("Per-subject on-device leaderboards with a podium and ranking.")
bullet("Background music and synthesized sound effects with user controls.")
bullet("Child-friendly privacy: no personal data is collected or sent anywhere.")

# ==================================================================
# 2. TECHNOLOGY STACK
# ==================================================================
h1("2. Technology Stack")
make_table(
    ["Layer", "Technology", "Purpose"],
    [
        ["Markup", "HTML5", "Page structure for all screens."],
        ["Styling", "CSS3 (custom style.css / Level.css)", "Theme, animations, responsive layout."],
        ["Logic", "Vanilla JavaScript (ES5/ES6)", "Game logic, state, progress, audio, UI."],
        ["CSS framework", "Bootstrap 5.3.0", "Grid, modals, dropdowns, responsive helpers."],
        ["Icons", "Font Awesome 6.4.0", "UI icons throughout the interface."],
        ["Fonts", "Google Fonts (DynaPuff, Irish Grover, Keania One)", "Playful display typography."],
        ["Effects", "canvas-confetti 1.9.3 (CDN)", "Celebration animation on level win."],
        ["Media", "Local image/audio assets + Unsplash fallback", "Card artwork and narration."],
        ["Storage", "Web Storage (localStorage / sessionStorage)", "Per-player progress, scores, settings."],
        ["Audio API", "Web Audio API (synthesized SFX)", "Click / wrong / game-over sounds."],
    ],
    widths=[1.4, 2.6, 2.8],
)

# ==================================================================
# 3. APPLICATION FLOW (PAGES)
# ==================================================================
h1("3. Application Flow")
p("The game is organized as a set of static HTML pages that hand off to one another "
  "through links and query-string parameters. The flow is:")
numbered("index.html \u2014 splash/loading screen with an animated progress bar.")
numbered("Start.html \u2014 login screen where the player enters a fun nickname.")
numbered("Home.html \u2014 main lobby with a welcome greeting and a \u201cPlay Now\u201d button.")
numbered("Subject.html \u2014 subject picker (EPP / Science / AP).")
numbered("Level.html \u2014 the level map for the chosen subject.")
numbered("Gameplay/*.html \u2014 the memory-matching board for a specific subject and level.")

h2("3.1 index.html — Loading Screen")
p("A full-screen splash shows the game logo riding an animated loading bar. When the "
  "bar reaches 100%, the page redirects to Start.html.")

h2("3.2 Start.html — Login")
p("The player enters a nickname (3\u201312 characters, letters, numbers, and spaces only). "
  "The nickname is stored in localStorage and used to namespace every piece of progress, "
  "so several children can share one device with separate saves. A \u201cFor Parents\u201d modal "
  "explains the privacy model.")

h2("3.3 Home.html — Lobby")
p("The lobby greets the player by nickname and offers a \u201cPlay Now\u201d button that leads "
  "to the subject picker. A reusable navigation bar (navbar.html) shows the player level, "
  "progress bar, star rating, and points, plus access to the leaderboard and settings.")

h2("3.4 Subject.html — Subject Picker")
p("Three full-width cards let the player choose a subject: Computer / EPP, Science, or "
  "Araling Panlipunan (AP). Selecting one opens that subject's level map.")

h2("3.5 Level.html — Level Map")
p("Each subject shows a whimsical island map with ten level nodes connected by a wavy "
  "golden path. Completed levels are marked, the best star rating is displayed on each "
  "node, and the player's character (JojoMa) stands on the current level. Locked levels "
  "must be unlocked with points after completing the previous level.")

# ==================================================================
# 4. SUBJECTS & CURRICULUM
# ==================================================================
h1("4. Subjects & Curriculum")
p("Every subject contains 10 levels, and each level is a themed set of cards. The "
  "topics follow an elementary-school learning progression.")

h2("4.1 Computer / EPP")
make_table(
    ["Level", "Topic"],
    [
        ["1", "Search Engines"],
        ["2", "Parts of a Search Engine Home Page"],
        ["3", "Web Browsers"],
        ["4", "Ways to Communicate Online"],
        ["5", "Netiquette: Do's and Don'ts"],
        ["6", "Domain Types"],
        ["7", "Parts of a Web Browser"],
        ["8", "Parts of an Email"],
        ["9", "Email Words"],
        ["10", "Parts of a Computer"],
    ],
    widths=[1.0, 4.0],
)

h2("4.2 Science")
make_table(
    ["Level", "Topic"],
    [
        ["1", "Matter"],
        ["2", "Solids"],
        ["3", "Liquids"],
        ["4", "Gases"],
        ["5", "Properties of Matter"],
        ["6", "Measuring Tools"],
        ["7", "Units of Measurement"],
        ["8", "Physical Changes"],
        ["9", "Digestive System"],
        ["10", "Respiratory System"],
    ],
    widths=[1.0, 4.0],
)

h2("4.3 Araling Panlipunan (AP)")
make_table(
    ["Level", "Topic"],
    [
        ["1", "Bodies of Water"],
        ["2", "Landforms"],
        ["3", "Southeast Asian Neighbors"],
        ["4", "Weather and Climate Terms"],
        ["5", "Climate and Seasons"],
        ["6", "Island Groups and Major Islands"],
        ["7", "More Landforms and Bodies of Water"],
        ["8", "Ancient Filipino Society and Life"],
        ["9", "Material Culture"],
        ["10", "Non-Material Culture"],
    ],
    widths=[1.0, 4.0],
)

# ==================================================================
# 5. GAME MECHANICS
# ==================================================================
h1("5. Game Mechanics")

h2("5.1 Core Gameplay")
p("Each level presents a shuffled grid of face-down cards. A card flip reveals either an "
  "image, a word, or a labeled picture; the goal is to find every matching pair. When a "
  "pair is matched, the game's helper character (JojoMa) shows a short typewriter-style "
  "description of the concept and, when available, plays a spoken narration clip.")

h2("5.2 Scaling Difficulty")
bullet("Card count grows with the level: total cards = 4 + (level \u00d7 2).")
bullet("Level 1 has 3 pairs (6 cards); Level 10 has 12 pairs (24 cards).")

h2("5.3 Timer")
p("Every level runs on a 120-second countdown. Running out of time ends the level in a "
  "\u201close\u201d state, showing the time, moves, and matched pairs. The player can pause at "
  "any time.")

h2("5.4 Star Rating")
p("A \u201cstar bar\u201d starts at 100% and drains as the player makes mistakes or spends time: "
  "each wrong match subtracts 4 points from the bar, and each second subtracts 0.2. The "
  "final bar value maps to a 1\u20133 star rating using thresholds at 5%, 30%, and 52%.")

h2("5.5 Points & Unlocks")
bullet("Base points for a level = 10 + (level \u00d7 5).")
bullet("More stars earn bonus points: each star above 1 adds 50% of the base.")
bullet("Points are only awarded the first time a level is completed.")
bullet("Locked levels cost a flat 15 points to unlock.")
bullet("A level can only be unlocked after the previous level is completed, preventing "
       "the player from skipping ahead.")

h2("5.6 Player Level")
p("The player's overall level is derived from total completed levels across all subjects: "
  "player level = floor(completed / 3) + 1, with a progress bar showing progress within "
  "the current level (3 completions per level-up).")

# ==================================================================
# 6. PLAYER & DATA SYSTEM
# ==================================================================
h1("6. Player & Data System")
p("All persistence uses the browser's Web Storage APIs. Progress keys are namespaced by "
  "the player's nickname, so each nickname has an independent save.")

h2("6.1 Storage Model")
bullet("Nickname is stored under the key mm_player_nickname.")
bullet("Every per-player key is prefixed mm_{nickname}_ (for example "
       "mm_jojo_unlocked_computer).")
bullet("Session-only flags (fresh login, just-unlocked level, music position) use "
       "sessionStorage.")

h2("6.2 Data Tracked Per Player")
make_table(
    ["Data", "Description"],
    [
        ["unlocked_{subject}", "List of levels unlocked for a subject (starts at [1])."],
        ["completed_{subject}", "List of levels completed for a subject."],
        ["stars_{subject}", "Best star rating per level."],
        ["pointsTotal", "Total points available to spend on unlocks."],
        ["gamesPlayed / totalMatches / bestTime / rewardsCount", "Lifetime player statistics."],
        ["Leaderboard entries", "Per-subject score records (nickname only)."],
    ],
    widths=[2.6, 3.6],
)

h2("6.3 Reset Progress")
p("The settings panel includes a \u201cReset Progress\u201d action that wipes only the current "
  "player's namespaced keys and leaderboard entries, leaving other players on the same "
  "device unaffected.")

# ==================================================================
# 7. LEADERBOARD
# ==================================================================
h1("7. Leaderboard")
p("Each subject has its own on-device leaderboard, shown in a modal with a podium for the "
  "top three and a ranked list below. There is no seeded/fake data \u2014 the board starts "
  "empty and only real wins are recorded.")
p("Ranking order:")
numbered("Highest level reached.")
numbered("Fastest completion time.")
numbered("Most total stars.")
numbered("Earliest recorded date.")
bullet("Display limit: 10 rows; storage keeps up to 50 per subject.", )
bullet("One row per nickname per subject; only the nickname is stored.", )

# ==================================================================
# 8. AUDIO & SOUND
# ==================================================================
h1("8. Audio & Sound")
h2("8.1 Background Music")
p("A background music module attempts to autoplay on load and, if blocked by the browser, "
  "retries on the first user interaction (click, key press, scroll, touch, etc.). Playback "
  "position is preserved across pages within a session, and the music can be toggled in "
  "settings.")

h2("8.2 Sound Effects")
p("Interface sound effects (click, wrong match, game over) are synthesized in real time "
  "with the Web Audio API \u2014 no audio files are required \u2014 so they work instantly and "
  "respond to the sound toggle in settings.")

h2("8.3 Card Narration")
p("When a card pair is matched, the game optionally plays a short spoken description "
  "using per-card audio assets, with graceful fallbacks if the file is missing.")

# ==================================================================
# 9. SETTINGS
# ==================================================================
h1("9. Settings")
make_table(
    ["Setting", "Options / Behavior"],
    [
        ["Sound Effects", "On / Off toggle."],
        ["Background Music", "On / Off toggle (live, cross-tab)."],
        ["Difficulty", "Easy / Normal / Hard."],
        ["Reset Progress", "Wipes current player's progress and leaderboard rows."],
    ],
    widths=[2.0, 4.2],
)

# ==================================================================
# 10. PRIVACY & SAFETY
# ==================================================================
h1("10. Privacy & Safety")
p("Match Monster is designed to be child-safe and privacy-respecting:")
bullet("Only a fun nickname is collected \u2014 never real names, emails, ages, or locations.")
bullet("All data (nickname, scores, progress) stays on the device; nothing is sent to a "
       "server or the internet.")
bullet("The login screen explicitly advises children to use a fun name instead of their "
       "real name.")
bullet("The project states compliance with the Philippine Data Privacy Act by not "
       "collecting personal information from children.")

# ==================================================================
# 11. FILE STRUCTURE
# ==================================================================
h1("11. File Structure")
make_table(
    ["Path", "Contents"],
    [
        ["index.html", "Loading / splash screen."],
        ["Start.html", "Nickname login + parent info modal."],
        ["Home.html", "Main lobby + navigation."],
        ["Subject.html", "Subject picker."],
        ["Level.html", "Level map for a subject."],
        ["navbar.html", "Shared navigation bar + leaderboard/settings modals."],
        ["script/", "15 JavaScript modules (player, levels, leaderboard, audio, gameplay, etc.)."],
        ["style/", "style.css and Level.css (theme + layout)."],
        ["Gameplay/", "Gameplay-computer.html, Gameplay-science.html, Gameplay-AP.html."],
        ["data/", "leaderboard.json (empty seed data)."],
        ["Assets/", "Images, icons, clouds, islands, topic pictures, and sound files."],
    ],
    widths=[1.8, 4.4],
)
p("The Assets/ directory holds 400+ media files, organized into subfolders for icons, "
  "clouds, islands, subjects, and per-topic/per-level card images and audio.", size=10,
  italic=True, color=GREY)

# ==================================================================
# 12. SUMMARY
# ==================================================================
h1("12. Summary")
p("Match Monster is a complete, self-contained educational game that demonstrates a "
  "thoughtful blend of game design and learning content. It pairs a familiar memory-match "
  "loop with a points-and-stars progression system, a multi-subject curriculum, a friendly "
  "monster mascot, and a strong privacy posture \u2014 all delivered as static web pages with "
  "no backend required.")

doc.save("Match_Monster_Documentation.docx")
print("Saved Match_Monster_Documentation.docx")
